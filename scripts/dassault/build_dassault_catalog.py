#!/usr/bin/env python3
"""Build the Dassault Systèmes SFDMU data plans from the portfolio workbook.

Reads the "Reference Portfolio" and "Industry Portfolio" sheets and writes the
CSVs for two plans:

  datasets/sfdmu/dassault/en-US/dassault-pcm      catalog structure
  datasets/sfdmu/dassault/en-US/dassault-pricing  price books and entries

Usage:
  python scripts/dassault/build_dassault_catalog.py "<path to workbook>.xlsx"

Requires openpyxl.
"""

import csv
import sys
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
PCM_DIR = REPO / "datasets/sfdmu/dassault/en-US/dassault-pcm"
PRICING_DIR = REPO / "datasets/sfdmu/dassault/en-US/dassault-pricing"

STANDARD_PRICEBOOK = "Standard Price Book"
PRORATION_POLICY = "Default Proration Policy"

# Workbook price column -> existing org selling model (name, type). Salesforce allows
# one selling model per type/term/unit and a term of 1 for Annual, so the org's
# standard models are reused. TBL2, TBL3 and TSC3 are priced through the Commitment
# attribute (COMMITMENT_PRICES); ELC is zero throughout the source and is not loaded.
SELLING_MODELS = OrderedDict(
    [
        ("PLC", ("One-Time", "OneTime")),
        ("ALC", ("Evergreen Annual", "Evergreen")),
        ("YLC", ("Term Annual", "TermDefined")),
        ("QLC", ("Term Based - Quarterly", "TermDefined")),
        ("PSC", ("One-Time", "OneTime")),
        ("ASC", ("Evergreen Annual", "Evergreen")),
        ("YSC", ("Term Annual", "TermDefined")),
    ]
)
DEFAULT_MODEL_PRIORITY = ["YLC", "YSC", "PLC", "PSC"]

PRICEBOOK_BY_SOURCE = {
    "EuroZone Price List": "EuroZone Price List",
    "USA Price List": "USA Price List",
}

RELEASE_LABELS = {
    "V5-6R2026": "V5-6R2026",
    "3DEXPERIENCER2026x": "3DEXPERIENCE R2026x",
    "3DEXPERIENCE R2026x": "3DEXPERIENCE R2026x",
}

CLASSIFICATIONS = OrderedDict(
    [
        ("DS-PC-V5-ONPREM", ("DELMIA V5 (On Premise)", "V5-6R2026", "On Premise")),
        ("DS-PC-3DX-ONPREM", ("3DEXPERIENCE (On Premise)", "3DEXPERIENCE R2026x", "On Premise")),
        ("DS-PC-3DX-CLOUD", ("3DEXPERIENCE (Public Cloud)", "3DEXPERIENCE R2026x", "Public Cloud")),
    ]
)

ATTR_CATEGORY = ("DS-ATTR-LICENSING", "Dassault Licensing", "Licensing and release characteristics of a Dassault Systèmes item")

ON_PREMISE_CLASSES = ("DS-PC-V5-ONPREM", "DS-PC-3DX-ONPREM")
CLOUD_CLASSES = ("DS-PC-3DX-CLOUD",)
ONE_YEAR = "1 Year"

# code, label, developer name, picklist name (None = text), read-only, required,
# classifications (None = all), price impacting
ATTRIBUTES = [
    ("DS_RELEASE", "Release", "DS_Release", "DS Release", False, True, None, False),
    ("DS_LICENSE_SCHEME", "License Scheme", "DS_License_Scheme", "DS License Scheme", True, False, None, False),
    ("DS_LICENSE_TYPE", "License Type", "DS_License_Type", "DS License Type", True, False, None, False),
    ("DS_DEPLOYMENT", "Deployment", "DS_Deployment", "DS Deployment", True, False, None, False),
    ("DS_REFERENCE_PORTFOLIO", "Reference Portfolio", "DS_Reference_Portfolio", None, True, False, None, False),
    ("DS_LICENSE_COMMITMENT", "Commitment", "DS_License_Commitment", "DS License Commitment", False, True,
     ON_PREMISE_CLASSES, True),
    ("DS_SUBSCRIPTION_COMMITMENT", "Commitment", "DS_Subscription_Commitment", "DS Subscription Commitment", False,
     True, CLOUD_CLASSES, True),
]
ATTR_DESCRIPTIONS = {
    "DS_RELEASE": "Software release delivered to the customer. Printed on the quote and invoice.",
    "DS_LICENSE_SCHEME": "Usage right granted by the product number (Named User, Concurrent, Add-On, ...).",
    "DS_LICENSE_TYPE": "Technical implementation of the item; drives key/media behaviour.",
    "DS_DEPLOYMENT": "On Premise or Public Cloud delivery.",
    "DS_REFERENCE_PORTFOLIO": "Reference portfolio of the item. Printed on the quote whatever catalog it was sold from.",
    "DS_LICENSE_COMMITMENT": "Term license commitment. 2 and 3 years apply the TBL2 / TBL3 price to the yearly license.",
    "DS_SUBSCRIPTION_COMMITMENT": "Cloud subscription commitment. 3 years applies the TSC3 price to the yearly subscription.",
}

# AttributeDefinition.Name must be unique; both commitments display as "Commitment".
ATTR_NAMES = {
    "DS_LICENSE_COMMITMENT": "Commitment (License)",
    "DS_SUBSCRIPTION_COMMITMENT": "Commitment (Subscription)",
}

PICKLISTS = OrderedDict(
    [
        ("DS Release", ["V5-6R2026", "3DEXPERIENCE R2026x"]),
        ("DS License Scheme", ["Named User", "Casual Named User", "Concurrent", "Add-On", "Credit Base"]),
        ("DS License Type", ["Named User", "Configuration", "Package", "Add-On"]),
        ("DS Deployment", ["On Premise", "Public Cloud"]),
        ("DS License Commitment", [ONE_YEAR, "2 Years", "3 Years"]),
        ("DS Subscription Commitment", [ONE_YEAR, "3 Years"]),
    ]
)
PICKLIST_CODE_PREFIX = {
    "DS Release": "DS-REL",
    "DS License Scheme": "DS-LS",
    "DS License Type": "DS-LT",
    "DS Deployment": "DS-DEP",
    "DS License Commitment": "DS-LC",
    "DS Subscription Commitment": "DS-SC",
}

# Multi-year price column -> (commitment attribute, value, years covered by the amount).
# TBL2 / TBL3 are read as the total for the whole term and TSC3 as a yearly price,
# because TBL2 and TBL3 exceed the yearly YLC price while TSC3 is below YSC.
# The override is applied to the yearly Term Annual price.
COMMITMENT_PRICES = OrderedDict(
    [
        ("TBL2", ("DS_LICENSE_COMMITMENT", "2 Years", 2)),
        ("TBL3", ("DS_LICENSE_COMMITMENT", "3 Years", 3)),
        ("TSC3", ("DS_SUBSCRIPTION_COMMITMENT", "3 Years", 1)),
    ]
)
COMMITMENT_SELLING_MODEL = "Term Annual"
ADJUSTMENT_EFFECTIVE_FROM = "2026-01-01T00:00:00.000+0000"
ADJUSTMENT_SCHEDULES = OrderedDict(
    [
        ("EUR", ("DS EuroZone Attribute Based Adjustment", "EuroZone Price List")),
        ("USD", ("DS USA Attribute Based Adjustment", "USA Price List")),
    ]
)

REF_CATALOG = ("DS-CAT-REF", "Dassault Systèmes Reference Portfolio",
               "Reference portfolios: every product belongs to exactly one. Portfolio > Theme > Discipline.")
IND_CATALOG = ("DS-CAT-IND", "Dassault Systèmes Industry Portfolio",
               "Industry solution experiences: Industry > Segment > Industry Solution Experience > Industry Process Experience.")


def slug(value):
    return "".join(ch if ch.isalnum() else "-" for ch in value.upper()).strip("-")


def picklist_code(picklist, value):
    return f"{PICKLIST_CODE_PREFIX[picklist]}-{slug(value)}"


def write_csv(directory, name, header, rows):
    directory.mkdir(parents=True, exist_ok=True)
    with open(directory / f"{name}.csv", "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(header)
        for row in rows:
            writer.writerow(["" if v is None else v for v in row])
    print(f"  {directory.name}/{name}.csv: {len(rows)} rows")


def fmt_price(value):
    return f"{value:.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else str(value)


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    ref_rows = list(wb["Reference Portfolio"].iter_rows(values_only=True))
    header = [h.strip() if isinstance(h, str) else h for h in ref_rows[3]]
    col = {h: i for i, h in enumerate(header)}
    price_cols = {code: col[f"{code} Amount"] for code in SELLING_MODELS}
    commitment_cols = {code: col[f"{code} Amount"] for code in COMMITMENT_PRICES}
    reference = []
    for r in ref_rows[4:]:
        if not r or not r[col["Product Number"]]:
            continue
        reference.append({
            "pricebook": r[col["Price List"]],
            "currency": r[col["Currency"]],
            "portfolio": r[col["Portfolio Name"]],
            "portfolio_id": r[col["Portfolio internal ID"]],
            "theme": r[col["Theme"]],
            "theme_id": r[col["Theme internal ID (Domain ID)"]],
            "discipline": r[col["Discipline"]],
            "discipline_id": r[col["Discipline internal ID (Group ID)"]],
            "trig": r[col["Item Trig"]],
            "sku": r[col["Product Number"]],
            "scheme": r[col["Licene Scheme"]],
            "name": r[col["Item Name"]],
            "license_type": r[col["License Type"]],
            "version": r[col["Version Name"]],
            "release": RELEASE_LABELS[r[col["Latest availlable release"]]],
            "prices": {code: r[i] for code, i in price_cols.items() if r[i] not in (None, 0)},
            "commitment_prices": {code: r[i] for code, i in commitment_cols.items() if r[i] not in (None, 0)},
        })

    ind_rows = list(wb["Industry Portfolio"].iter_rows(values_only=True))
    ih = {h: i for i, h in enumerate(ind_rows[0])}
    industry = []
    for r in ind_rows[1:]:
        if not r or not r[ih["Product Number"]]:
            continue
        industry.append({k: r[ih[k]] for k in ih})
    return reference, industry


def build(path):
    reference, industry = read_workbook(path)

    industry_names = {r["Product Number"]: r["Item Name"] for r in industry}
    by_sku = OrderedDict()
    for row in reference:
        by_sku.setdefault(row["sku"], []).append(row)

    products = OrderedDict()
    for sku, rows in by_sku.items():
        first = rows[0]
        names = list(OrderedDict.fromkeys(r["name"] for r in rows))
        name = industry_names.get(sku) if industry_names.get(sku) in names else names[0]
        deployment = "Public Cloud" if first["portfolio"].startswith("Public Cloud") else "On Premise"
        if first["version"] == "Version 5":
            classification = "DS-PC-V5-ONPREM"
        else:
            classification = "DS-PC-3DX-CLOUD" if deployment == "Public Cloud" else "DS-PC-3DX-ONPREM"
        aliases = [n for n in names if n != name]
        description = (f"{first['portfolio']} > {first['theme']} > {first['discipline']}. "
                       f"{first['scheme']} license, {first['license_type']}, {first['version']}.")
        if aliases:
            description += " Also sold as: " + "; ".join(aliases) + "."
        products[sku] = {
            "sku": sku, "name": name, "trig": first["trig"], "family": first["portfolio"],
            "classification": classification, "description": description,
            "attrs": {
                "DS_RELEASE": first["release"], "DS_LICENSE_SCHEME": first["scheme"],
                "DS_LICENSE_TYPE": first["license_type"], "DS_DEPLOYMENT": deployment,
                "DS_REFERENCE_PORTFOLIO": first["portfolio"],
                "DS_LICENSE_COMMITMENT": ONE_YEAR, "DS_SUBSCRIPTION_COMMITMENT": ONE_YEAR,
            },
            "category": f"DS-{first['discipline_id']}",
            "models": [c for c in SELLING_MODELS if c in first["prices"]],
            "has_commitment": bool(first["commitment_prices"]),
            "priced": True,
        }

    for row in industry:
        sku = row["Product Number"]
        if sku in products:
            continue
        deployment = row["Cloud/Premise ?"]
        products[sku] = {
            "sku": sku, "name": row["Item Name"], "trig": row["Item Trig"], "family": row["Portfolio Name"],
            "classification": "DS-PC-3DX-CLOUD" if deployment == "Public Cloud" else "DS-PC-3DX-ONPREM",
            "description": (f"{row['Portfolio Name']} ({deployment}). Listed in the Industry Portfolio only; "
                            "no reference-portfolio price in the source workbook."),
            "attrs": {
                "DS_RELEASE": RELEASE_LABELS.get(row["Item Release Name"], row["Item Release Name"]),
                "DS_LICENSE_SCHEME": None, "DS_LICENSE_TYPE": None, "DS_DEPLOYMENT": deployment,
                "DS_REFERENCE_PORTFOLIO": row["Portfolio Name"],
                "DS_LICENSE_COMMITMENT": ONE_YEAR, "DS_SUBSCRIPTION_COMMITMENT": ONE_YEAR,
            },
            "category": None, "models": [], "has_commitment": False, "priced": False,
        }

    write_pcm(reference, industry, products)
    write_pricing(reference, products)
    unpriced = sum(1 for p in products.values() if not p["priced"])
    print(f"\n{len(products)} products ({len(products) - unpriced} priced, {unpriced} industry-only without price)")


def write_pcm(reference, industry, products):
    print("dassault-pcm")
    write_csv(PCM_DIR, "AttributePicklist", ["Code", "DataType", "Description", "Name", "Status"],
              [[name, "Text", f"{name[3:]} values for Dassault Systèmes items", name, "Active"] for name in PICKLISTS])

    values = []
    for name, options in PICKLISTS.items():
        for seq, value in enumerate(options, start=1):
            values.append([None, picklist_code(name, value), value, "false", value, name, seq, "Active", value])
    write_csv(PCM_DIR, "AttributePicklistValue",
              ["Abbreviation", "Code", "DisplayValue", "IsDefault", "Name", "Picklist.Name", "Sequence", "Status", "Value"],
              values)

    write_csv(PCM_DIR, "AttributeDefinition",
              ["Code", "DataType", "DefaultHelpText", "DefaultValue", "Description", "DeveloperName", "IsActive",
               "IsRequired", "Label", "Name", "Picklist.Name", "SourceSystemIdentifier", "ValueDescription"],
              [[code, "Picklist" if pl else "Text", None, None, ATTR_DESCRIPTIONS[code], dev, "true",
                "false", label, ATTR_NAMES.get(code, label), pl, None, None]
               for code, label, dev, pl, *_ in ATTRIBUTES])

    cat_code, cat_name, cat_desc = ATTR_CATEGORY
    write_csv(PCM_DIR, "AttributeCategory", ["Code", "Description", "Name"], [[cat_code, cat_desc, cat_name]])
    write_csv(PCM_DIR, "AttributeCategoryAttribute",
              ["$$AttributeCategory.Code$AttributeDefinition.Code", "AttributeCategory.Code", "AttributeDefinition.Code"],
              [[f"{cat_code};{a[0]}", cat_code, a[0]] for a in ATTRIBUTES])

    write_csv(PCM_DIR, "ProductClassification", ["Code", "Name", "Status"],
              [[code, name, "Active"] for code, (name, _rel, _dep) in CLASSIFICATIONS.items()])

    # A visible read-only attribute must carry a default, so attributes whose value
    # differs per product are only locked at product level.
    class_attr_rows = []
    for pc_code, (_name, release, deployment) in CLASSIFICATIONS.items():
        defaults = {"DS_RELEASE": release, "DS_DEPLOYMENT": deployment,
                    "DS_LICENSE_COMMITMENT": ONE_YEAR, "DS_SUBSCRIPTION_COMMITMENT": ONE_YEAR}
        for seq, (code, label, _dev, pl, read_only, required, classes, price_impacting) in enumerate(ATTRIBUTES, 1):
            if classes and pc_code not in classes:
                continue
            default = defaults.get(code)
            class_attr_rows.append([cat_code, code, None, default, None, "ComboBox" if pl else "Text",
                                    "false", str(price_impacting).lower(), str(read_only and default is not None).lower(),
                                    str(required).lower(), classification_attr_name(pc_code, label), pc_code, seq,
                                    "Active"])
    write_csv(PCM_DIR, "ProductClassificationAttr",
              ["AttributeCategory.Code", "AttributeDefinition.Code", "AttributeNameOverride", "DefaultValue",
               "Description", "DisplayType", "IsHidden", "IsPriceImpacting", "IsReadOnly", "IsRequired", "Name",
               "ProductClassification.Code", "Sequence", "Status"], class_attr_rows)

    write_csv(PCM_DIR, "Product2",
              ["BasedOn.Code", "ConfigureDuringSale", "Description", "Family", "IsActive", "IsAssetizable",
               "IsSoldOnlyWithOtherProds", "Name", "ProductCode", "StockKeepingUnit", "Type"],
              [[p["classification"], "Allowed", p["description"], p["family"], "true", "true", "false",
                p["name"], p["trig"], p["sku"], None] for p in products.values()])

    pad_rows = []
    for p in products.values():
        for seq, (code, label, _dev, pl, read_only, required, classes, price_impacting) in enumerate(ATTRIBUTES, 1):
            if classes and p["classification"] not in classes:
                continue
            value = p["attrs"][code]
            hidden = value is None or (price_impacting and not p["has_commitment"])
            pad_rows.append([f"{code};{p['sku']}", cat_code, code, None, value, None,
                             "ComboBox" if pl else "Text", str(hidden).lower(), str(price_impacting).lower(),
                             str(read_only and value is not None).lower(), str(required and not hidden).lower(),
                             label, p["sku"],
                             classification_attr_name(p["classification"], label), seq, "Active"])
    write_csv(PCM_DIR, "ProductAttributeDefinition",
              ["$$AttributeDefinition.Code$Product2.StockKeepingUnit", "AttributeCategory.Code",
               "AttributeDefinition.Code", "AttributeNameOverride", "DefaultValue", "Description", "DisplayType",
               "IsHidden", "IsPriceImpacting", "IsReadOnly", "IsRequired", "Name", "Product2.StockKeepingUnit",
               "ProductClassificationAttribute.Name", "Sequence", "Status"], pad_rows)

    write_csv(PCM_DIR, "ProductSellingModel", ["Name", "SellingModelType"], selling_model_rows())
    write_csv(PCM_DIR, "ProrationPolicy", ["Name"], [[PRORATION_POLICY]])

    psmo_rows = []
    for p in products.values():
        default = next((c for c in DEFAULT_MODEL_PRIORITY if c in p["models"]), p["models"][0] if p["models"] else None)
        for code in p["models"]:
            name, smtype = SELLING_MODELS[code]
            psmo_rows.append([f"{p['sku']};{name}", str(code == default).lower(), p["sku"], name,
                              None if smtype == "OneTime" else PRORATION_POLICY])
    write_csv(PCM_DIR, "ProductSellingModelOption",
              ["$$Product2.StockKeepingUnit$ProductSellingModel.Name", "IsDefault", "Product2.StockKeepingUnit",
               "ProductSellingModel.Name", "ProrationPolicy.Name"], psmo_rows)

    write_csv(PCM_DIR, "ProductCatalog",
              ["CatalogType", "Code", "Description", "EffectiveEndDate", "EffectiveStartDate", "Name"],
              [["Sales", code, desc, None, None, name] for code, name, desc in (REF_CATALOG, IND_CATALOG)])

    categories = OrderedDict()
    for r in reference:
        categories.setdefault(f"DS-{r['portfolio_id']}", (REF_CATALOG[0], r["portfolio"], None))
        categories.setdefault(f"DS-{r['theme_id']}", (REF_CATALOG[0], r["theme"], f"DS-{r['portfolio_id']}"))
        categories.setdefault(f"DS-{r['discipline_id']}", (REF_CATALOG[0], r["discipline"], f"DS-{r['theme_id']}"))
    for r in industry:
        industry_code = "DS-IND-" + slug(r["Industry"].split(" ", 1)[0])
        segment_code = "DS-SEG-" + slug(r["Segment"].split(" ", 1)[0])
        categories.setdefault(industry_code, (IND_CATALOG[0], r["Industry"].split(" ", 1)[1], None))
        categories.setdefault(segment_code, (IND_CATALOG[0], r["Segment"].split(" ", 1)[1], industry_code))
        categories.setdefault(f"DS-{r['ISE ID']}",
                              (IND_CATALOG[0], f"{r['ISE Name']} ({r['Cloud/Premise ?']})", segment_code))
        categories.setdefault(f"DS-{r['IPE ID']}", (IND_CATALOG[0], r["IPE Name"], f"DS-{r['ISE ID']}"))
    sort_order = Counter()
    category_rows = []
    for code, (catalog, name, parent) in categories.items():
        sort_order[(catalog, parent)] += 10
        category_rows.append([catalog, code, None, "true", name, parent, sort_order[(catalog, parent)]])
    write_csv(PCM_DIR, "ProductCategory",
              ["Catalog.Code", "Code", "Description", "IsNavigational", "Name", "ParentCategory.Code", "SortOrder"],
              category_rows)

    pcp_rows = []
    for p in products.values():
        if p["category"]:
            pcp_rows.append([f"{p['category']};{p['sku']}", REF_CATALOG[0], "true", p["category"], p["sku"]])
    for r in industry:
        category = f"DS-{r['IPE ID']}"
        pcp_rows.append([f"{category};{r['Product Number']}", IND_CATALOG[0],
                         str(not products[r["Product Number"]]["category"]).lower(), category, r["Product Number"]])
    write_csv(PCM_DIR, "ProductCategoryProduct",
              ["$$ProductCategory.Code$Product.StockKeepingUnit", "Catalog.Code", "IsPrimaryCategory",
               "ProductCategory.Code", "Product.StockKeepingUnit"], pcp_rows)


def classification_attr_name(pc_code, label):
    return f"{CLASSIFICATIONS[pc_code][0]} - {label}"


def selling_model_rows():
    return [list(model) for model in OrderedDict.fromkeys(SELLING_MODELS.values())]


def write_pricing(reference, products):
    """Pass 1 loads the price books, Standard Price Book entries and commitment
    adjustments; pass 2 (object-set-2) loads the EuroZone / USA entries, which
    Salesforce rejects until a standard price exists."""
    print("dassault-pricing")
    pass2_dir = PRICING_DIR / "objectset_source" / "object-set-2"
    lookups = {
        "Pricebook2": (["Description", "IsActive", "IsStandard", "Name"],
                       [[None, "true", "true", STANDARD_PRICEBOOK],
                        ["Dassault Systèmes price list for EuroZone customers (EUR)", "true", "false",
                         "EuroZone Price List"],
                        ["Dassault Systèmes price list for USA customers (USD)", "true", "false", "USA Price List"]]),
        "ProductSellingModel": (["Name", "SellingModelType"], selling_model_rows()),
        "Product2": (["StockKeepingUnit"], [[sku] for sku, p in products.items() if p["priced"]]),
    }
    for directory in (PRICING_DIR, pass2_dir):
        for name, (header, rows) in lookups.items():
            write_csv(directory, name, header, rows)

    prices = OrderedDict()
    for r in reference:
        for code, amount in r["prices"].items():
            prices.setdefault((r["sku"], code, r["currency"]), (r["pricebook"], amount))

    header = ["$$Pricebook2.Name$Product2.StockKeepingUnit$ProductSellingModel.Name$CurrencyIsoCode",
              "CurrencyIsoCode", "IsActive", "IsDerived", "Pricebook2.Name", "Product2.StockKeepingUnit",
              "ProductSellingModel.Name", "UnitPrice", "UseStandardPrice"]
    standard_rows, list_rows = [], []
    for (sku, code, currency), (source_pricebook, amount) in prices.items():
        model = SELLING_MODELS[code][0]
        for pricebook, target in ((STANDARD_PRICEBOOK, standard_rows), (PRICEBOOK_BY_SOURCE[source_pricebook], list_rows)):
            target.append([f"{pricebook};{sku};{model};{currency}", currency, "true", "false", pricebook, sku,
                           model, fmt_price(amount), "false"])
    write_csv(PRICING_DIR, "PricebookEntry", header, standard_rows)
    write_csv(pass2_dir, "PricebookEntry", header, list_rows)
    write_commitment_adjustments(reference, PRICING_DIR)


def write_commitment_adjustments(reference, directory):
    """Attribute-based price overrides for the multi-year commitments (pass 1)."""
    write_csv(directory, "AttributeDefinition", ["Code"],
              [[code] for code in sorted({attribute for attribute, *_ in COMMITMENT_PRICES.values()})])
    write_csv(directory, "PriceAdjustmentSchedule",
              ["AdjustmentMethod", "CurrencyIsoCode", "Description", "EffectiveFrom", "IsActive", "Name",
               "Pricebook2.Name", "ScheduleType"],
              [["Range", currency, f"Dassault Systèmes multi-year commitment prices for the {pricebook}",
                ADJUSTMENT_EFFECTIVE_FROM, "true", name, pricebook, "Attribute"]
               for currency, (name, pricebook) in ADJUSTMENT_SCHEDULES.items()])

    conditions, adjustments = OrderedDict(), OrderedDict()
    for r in reference:
        for code, amount in r["commitment_prices"].items():
            attribute, value, years = COMMITMENT_PRICES[code]
            rule = f"DS {r['sku']} Commitment {value}"
            conditions.setdefault(rule, [f"{rule};{attribute};{r['sku']}", rule, attribute, "equals",
                                         f"{rule} condition", r["sku"], value])
            schedule = ADJUSTMENT_SCHEDULES[r["currency"]][0]
            key = f"{rule};{schedule};{r['sku']};{COMMITMENT_SELLING_MODEL};{r['currency']}"
            adjustments.setdefault(key, [key, "Override", fmt_price(round(amount / years, 2)), rule, r["currency"],
                                         ADJUSTMENT_EFFECTIVE_FROM, f"{rule} {r['currency']}", schedule, r["sku"],
                                         COMMITMENT_SELLING_MODEL])
    write_csv(directory, "AttributeBasedAdjRule", ["Name"], [[rule] for rule in conditions])
    write_csv(directory, "AttributeAdjustmentCondition",
              ["$$AttributeBasedAdjRule.Name$AttributeDefinition.Code$Product.StockKeepingUnit",
               "AttributeBasedAdjRule.Name", "AttributeDefinition.Code", "Operator", "Name",
               "Product.StockKeepingUnit", "StringValue"], list(conditions.values()))
    write_csv(directory, "AttributeBasedAdjustment",
              ["$$AttributeBasedAdjRule.Name$PriceAdjustmentSchedule.Name$Product.StockKeepingUnit"
               "$ProductSellingModel.Name$CurrencyIsoCode",
               "AdjustmentType", "AdjustmentValue", "AttributeBasedAdjRule.Name", "CurrencyIsoCode", "EffectiveFrom",
               "Name", "PriceAdjustmentSchedule.Name", "Product.StockKeepingUnit", "ProductSellingModel.Name"],
              list(adjustments.values()))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    build(sys.argv[1])
