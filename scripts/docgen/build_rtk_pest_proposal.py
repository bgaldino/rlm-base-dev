"""Build the Rentokil pest proposal DocGen template (docx).

Usage:
    python scripts/docgen/build_rtk_pest_proposal.py \
        scripts/docgen/assets/rentokil_logo_white.png \
        unpackaged/post_pest_docgen/documentTemplates/RLM_RTK_PestProposal_1.dt

Tokens must match the RLMRTKProposalTransform output and the template's tokenList.
"""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RED = "EC1D23"
INK = "071017"
GREY = "5A6470"
TINT = "FDECEC"
RULE = "D9DDE2"
LOGO_WHITE = sys.argv[1]
OUT = sys.argv[2]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.6)
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.6)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.3)
CONTENT_W = 7.3

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def borders(table, color=None, inside_h=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if color and (edge == "insideH" and inside_h or edge == "bottom"):
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def widths(table, inches):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, inches):
            cell.width = Inches(w)


def run(p, text, size=None, bold=False, color=None, italic=False, caps=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.all_caps = caps
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(container, text="", size=None, bold=False, color=None, align=None, before=0, after=0, caps=False, italic=False):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, color, italic, caps)
    return p


def cell_text(cell, text, size=None, bold=False, color=None, align=None, caps=False):
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run(p, text, size, bold, color, caps=caps)
    return p


def section_title(text, before=14):
    p = para(doc, before=before, after=4)
    run(p, text, 11, True, RED, caps=True)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), RED)
    bdr.append(bottom)
    pPr.append(bdr)
    return p


# ---------- header: red band with logo + contact ----------
header = sec.header
hp = header.paragraphs[0]
ht = header.add_table(rows=1, cols=2, width=Inches(CONTENT_W))
borders(ht)
cell_margins(ht, 140, 140, 180, 180)
widths(ht, [3.0, 4.3])
for c in ht.rows[0].cells:
    shade(c, RED)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
ht.cell(0, 0).paragraphs[0].add_run().add_picture(LOGO_WHITE, width=Inches(1.9))
rc = ht.cell(0, 1)
cell_text(rc, "Pest Protection Proposal", 14, True, "FFFFFF", WD_ALIGN_PARAGRAPH.RIGHT)
p = rc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p, "Proposal {{QuoteNumber}}   |   Prepared {{CreatedDate}}", 8.5, False, "FFFFFF")
hp._p.getparent().remove(hp._p)
header.add_paragraph()

# ---------- footer ----------
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Rentokil Pest Control", 8, True, RED)
run(fp, "   |   Proposal {{QuoteNumber}}   |   Questions? Contact {{SalesRep}} at {{SalesRepEmail}}", 8, False, GREY)

# ---------- proposal meta ----------
p = para(doc, before=2)
run(p, "Prepared for ", 18, False, GREY)
run(p, "{{AccountName}}", 18, True, INK)
p = para(doc, before=2, after=8)
run(p, "Proposal {{QuoteNumber}}   ·   Prepared {{CreatedDate}}   ·   Pricing valid for 30 days", 9, False, GREY)

info = doc.add_table(rows=2, cols=3)
borders(info)
cell_margins(info, 50, 50, 120, 120)
widths(info, [2.45, 2.45, 2.4])
for i, label in enumerate(["SERVICE ADDRESS", "BILLING ADDRESS", "YOUR PEST SPECIALIST"]):
    c = info.cell(0, i)
    shade(c, TINT)
    cell_text(c, label, 7.5, True, RED)
blocks = [
    ["{{ServiceStreet}}", "{{ServiceCity}}, {{ServiceState}} {{ServicePostalCode}}"],
    ["{{BillingStreet}}", "{{BillingCity}}, {{BillingState}} {{BillingPostalCode}}", "{{CustomerPhone}}"],
    ["{{SalesRep}}", "{{SalesRepEmail}}"],
]
for i, lines in enumerate(blocks):
    c = info.cell(1, i)
    shade(c, TINT)
    first = True
    for line in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        run(p, line, 9.5, first, INK)
        first = False

# ---------- protection plans ----------
section_title("Your Protection Plan")
para(doc, "Ongoing service plans that keep your home protected all year. Billed monthly; cancel any time with 30 days' notice.",
     8.5, color=GREY, after=6)
plans = doc.add_table(rows=3, cols=3)
borders(plans, RULE, inside_h=True)
cell_margins(plans, 70, 70, 120, 120)
widths(plans, [2.2, 3.9, 1.2])
for i, h in enumerate(["PLAN", "WHAT'S INCLUDED", "PER MONTH"]):
    c = plans.cell(0, i)
    shade(c, INK)
    cell_text(c, h, 8, True, "FFFFFF", WD_ALIGN_PARAGRAPH.RIGHT if i == 2 else None)
cell_text(plans.cell(1, 0), "{{#Plan}}{{Name}}", 10, True, INK)
cell_text(plans.cell(1, 1), "{{Description}}", 9, False, GREY)
cell_text(plans.cell(1, 2), "{{MonthlyPrice}}{{/Plan}}", 10, True, INK, WD_ALIGN_PARAGRAPH.RIGHT)
tot = plans.rows[2].cells
merged = tot[0].merge(tot[1])
for c in (merged, tot[2]):
    shade(c, TINT)
cell_text(merged, "Monthly service total", 10, True, INK)
cell_text(tot[2], "{{MonthlyTotal}}", 11, True, RED, WD_ALIGN_PARAGRAPH.RIGHT)

# ---------- one-time services ----------
section_title("One-Time Services")
para(doc, "Performed at your first visit and billed once the treatment is complete.", 8.5, color=GREY, after=6)
one = doc.add_table(rows=3, cols=3)
borders(one, RULE, inside_h=True)
cell_margins(one, 70, 70, 120, 120)
widths(one, [2.2, 3.9, 1.2])
for i, h in enumerate(["SERVICE", "DETAILS", "PRICE"]):
    c = one.cell(0, i)
    shade(c, INK)
    cell_text(c, h, 8, True, "FFFFFF", WD_ALIGN_PARAGRAPH.RIGHT if i == 2 else None)
cell_text(one.cell(1, 0), "{{#OneTime}}{{Name}}", 10, True, INK)
cell_text(one.cell(1, 1), "{{Description}}", 9, False, GREY)
cell_text(one.cell(1, 2), "{{Price}}{{/OneTime}}", 10, True, INK, WD_ALIGN_PARAGRAPH.RIGHT)
tot = one.rows[2].cells
merged = tot[0].merge(tot[1])
for c in (merged, tot[2]):
    shade(c, TINT)
cell_text(merged, "One-time total", 10, True, INK)
cell_text(tot[2], "{{OneTimeTotal}}", 11, True, RED, WD_ALIGN_PARAGRAPH.RIGHT)

# ---------- summary band ----------
para(doc, after=8)
band = doc.add_table(rows=1, cols=3)
borders(band)
cell_margins(band, 110, 110, 160, 160)
widths(band, [2.45, 2.45, 2.4])
for i, (label, value) in enumerate([("MONTHLY SERVICE", "{{MonthlyTotal}} / mo"),
                                    ("DUE AFTER FIRST VISIT", "{{OneTimeTotal}}"),
                                    ("FIRST MONTH TOTAL", "{{GrandTotal}}")]):
    c = band.cell(0, i)
    shade(c, RED)
    cell_text(c, label, 7.5, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, value, 15, True, "FFFFFF")

# ---------- visits ----------
section_title("What Happens Next", before=16)
steps = [
    ("1", "Initial Service visit",
     "Once you accept, we schedule your first visit automatically. Your technician inspects the property, "
     "treats active activity, and seals common entry points."),
    ("2", "Routine protection visits",
     "Your plan visits follow the schedule described above, such as quarterly general pest service and "
     "seasonal mosquito treatments from April through October."),
    ("3", "Free re-service between visits",
     "If covered pests come back between scheduled visits, call us and we will return at no extra charge."),
    ("4", "A report after every visit",
     "You receive a service report listing what we found, what we treated, and any recommendations for your home."),
]
st = doc.add_table(rows=len(steps), cols=2)
borders(st)
cell_margins(st, 60, 60, 100, 100)
widths(st, [0.45, 6.85])
for r, (num, title, body) in enumerate(steps):
    n = st.cell(r, 0)
    shade(n, RED)
    n.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_text(n, num, 13, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    c = st.cell(r, 1)
    cell_text(c, title, 10, True, INK)
    p = c.add_paragraph()
    run(p, body, 9, False, GREY)

# ---------- promise ----------
section_title("Our Promise to You", before=16)
for line in [
    "Licensed, trained technicians who arrive in a marked vehicle and uniform.",
    "Treatments applied to label directions, with guidance on keeping children and pets safe.",
    "If you are not satisfied with a visit, tell us and we will make it right.",
]:
    p = para(doc, after=2)
    run(p, "■  ", 8, True, RED)
    run(p, line, 9, False, INK)

# ---------- terms ----------
section_title("Service Terms", before=16)
terms = [
    "Service plans renew monthly until cancelled by either party with 30 days' written notice.",
    "One-time services are invoiced after the treatment is completed. Monthly plans are billed in advance.",
    "Prices in this proposal are valid for 30 days from {{CreatedDate}} and assume the property details captured during inspection.",
    "Please provide access to the areas being treated and follow your technician's re-entry guidance after each visit.",
    "Structural repairs, moisture correction, and wildlife removal are not included unless listed above.",
]
for t in terms:
    p = para(doc, after=2)
    run(p, "•  ", 8.5, True, GREY)
    run(p, t, 8.5, False, GREY)

# ---------- signatures ----------
section_title("Acceptance", before=16)
para(doc, "By signing below, you accept this proposal and authorize Rentokil to begin service at the address above.",
     8.5, color=GREY, after=10)
sig = doc.add_table(rows=2, cols=2)
borders(sig)
cell_margins(sig, 40, 40, 0, 200)
widths(sig, [3.65, 3.65])
for i, who in enumerate(["Customer signature", "Rentokil representative"]):
    c = sig.cell(0, i)
    cell_text(c, " ", 18)
    p = c.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:color"), INK)
    bdr.append(b)
    pPr.append(bdr)
    c2 = sig.cell(1, i)
    cell_text(c2, who, 8, True, INK)
    p = c2.add_paragraph()
    run(p, "{{AccountName}}" if i == 0 else "{{SalesRep}}", 8.5, False, GREY)
    p = c2.add_paragraph()
    run(p, "Date: ____________________", 8.5, False, GREY)

doc.save(OUT)
print("saved", OUT)
